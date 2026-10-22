"""Bounded, read-only document previews. No public URLs or third-party viewers.

Office macros, formulas, embedded objects and external links are never executed.
The original remains available for full-fidelity editing in a desktop application.
"""
from __future__ import annotations

import csv
import io
import pathlib
import re
import zipfile
from datetime import date, datetime
from itertools import islice
from contextlib import contextmanager
from threading import BoundedSemaphore

from xml.parsers import expat
from fastapi import HTTPException

MAX_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_XML_METADATA_BYTES = 2 * 1024 * 1024
MAX_XML_ROW_BYTES = 1024 * 1024
MAX_XML_NODES = 500_000
MAX_CSV_RECORD_CHARS = 1_000_000
MAX_ROWS = 500
MAX_COLUMNS = 60
MAX_CELL_CHARS = 2000
MAX_TEXT_CHARS = 300_000
_preview_slots = BoundedSemaphore(2)
SUPPORTED = {"xlsx", "xlsm", "xls", "csv", "tsv", "docx", "pptx", "txt", "md", "json", "log"}



@contextmanager
def preview_slot():
    # Bound downloads plus parser memory per API process, not just per document.
    if not _preview_slots.acquire(blocking=False):
        raise HTTPException(429, "Сейчас открывается много документов. Повторите просмотр через несколько секунд.", headers={"Retry-After": "2"})
    try:
        yield
    finally:
        _preview_slots.release()


def _check_zip(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len({entry.filename for entry in entries}) != len(entries):
            raise HTTPException(422, "Документ содержит повторяющиеся части. Скачайте оригинал.")
        if len(entries) > 5000 or sum(x.file_size for x in entries) > MAX_EXPANDED_BYTES:
            raise HTTPException(413, "Документ слишком большой для предпросмотра. Скачайте оригинал.")
        if any(x.flag_bits & 1 for x in entries):
            raise HTTPException(422, "Документ защищён паролем. Откройте оригинал на компьютере.")
        if any(x.file_size > 50 * 1024 * 1024 or (x.file_size > 10 * 1024 * 1024 and x.file_size / max(x.compress_size, 1) > 300) for x in entries):
            raise HTTPException(413, "Структура документа превышает безопасные лимиты просмотра.")



class _PreviewComplete(Exception):
    pass


def _structure_limit() -> HTTPException:
    return HTTPException(413, "Структура документа слишком сложная для безопасного просмотра. Скачайте оригинал.")


def _walk_xml(stream, *, start=None, end=None, text=None, budget=None) -> None:
    """SAX-style bounded parsing: no full XML tree or archive.read() allocation."""
    parser = expat.ParserCreate(namespace_separator="}")
    budget = budget if budget is not None else {"nodes": 0}
    depth, fed = 0, 0

    def on_start(name, attributes):
        nonlocal depth
        depth += 1
        budget["nodes"] += 1
        if depth > 64 or budget["nodes"] > MAX_XML_NODES or len(attributes) > 64 or any(len(v) > 8192 for v in attributes.values()):
            raise _structure_limit()
        if start:
            start(name, attributes, parser.CurrentByteIndex)

    def on_end(name):
        nonlocal depth
        if end:
            end(name, parser.CurrentByteIndex)
        depth -= 1

    def forbidden(*_):
        raise HTTPException(422, "XML с DTD или внешними сущностями не поддерживается. Скачайте оригинал.")

    parser.StartElementHandler, parser.EndElementHandler = on_start, on_end
    parser.CharacterDataHandler = text or (lambda _: None)
    parser.StartDoctypeDeclHandler = forbidden
    parser.EntityDeclHandler = forbidden
    parser.ExternalEntityRefHandler = forbidden
    while chunk := stream.read(8192):
        fed += len(chunk)
        parser.Parse(chunk, False)
        # Expat buffers unfinished tokens. Bound a single giant attribute/tag,
        # not just decompressed file size or the eventual output text.
        if fed - parser.CurrentByteIndex > 256 * 1024:
            raise _structure_limit()
    parser.Parse(b"", True)


def _check_workbook_structure(data: bytes) -> dict[str, tuple[int, int]]:
    """openpyxl's read-only mode still loads styles/shared strings and a whole row.

    Preflight those allocations with streaming XML before invoking the library.
    This retains its date/formula/coordinate handling without unbounded trees.
    """
    budget = {"nodes": 0}
    dimensions = {}
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        metadata_bytes = 0
        for info in archive.infolist():
            if not info.filename.endswith((".xml", ".rels")):
                continue
            worksheet = bool(re.fullmatch(r"xl/worksheets/[^/]+\.xml", info.filename))
            if not worksheet:
                metadata_bytes += info.file_size
                if info.file_size > MAX_XML_METADATA_BYTES or metadata_bytes > 8 * MAX_XML_METADATA_BYTES:
                    raise _structure_limit()
            row_start, row_nodes, sheets, metadata_nodes = None, 0, 0, 0
            text_lengths = []
            source_rows = source_columns = cell_index = 0

            def start(name, attrs, position):
                nonlocal row_start, row_nodes, sheets, metadata_nodes, source_rows, source_columns, cell_index
                metadata_nodes += 1
                if not worksheet and metadata_nodes > 50000:
                    raise _structure_limit()
                text_lengths.append(0)
                local = name.rsplit("}", 1)[-1]
                # A part named .bin but declared as XML must not bypass preflight.
                if local == "Override" and "+xml" in attrs.get("ContentType", "") and not attrs.get("PartName", "").endswith(".xml"):
                    raise _structure_limit()
                if local == "Relationship" and attrs.get("Type", "").endswith(("/worksheet", "/sharedStrings", "/styles")) and not attrs.get("Target", "").endswith(".xml"):
                    raise _structure_limit()
                if worksheet and local == "row":
                    row_start, row_nodes, cell_index = position, 0, 0
                    source_rows = max(source_rows + 1 if "r" not in attrs else int(attrs["r"]), source_rows)
                if worksheet and local == "c":
                    cell_index += 1
                    coordinate = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", attrs.get("r", ""))
                    if coordinate:
                        cell_index = 0
                        for letter in coordinate.group(1):
                            cell_index = cell_index * 26 + ord(letter) - 64
                        source_rows = max(source_rows, int(coordinate.group(2)))
                    source_columns = max(source_columns, cell_index)
                if row_start is not None:
                    row_nodes += 1
                    if row_nodes > 20000 or position - row_start > MAX_XML_ROW_BYTES:
                        raise _structure_limit()
                if info.filename == "xl/workbook.xml" and local == "sheet":
                    sheets += 1
                    if sheets > 200:
                        raise _structure_limit()

            def end(name, position):
                nonlocal row_start
                text_lengths.pop()
                if row_start is not None and position - row_start > MAX_XML_ROW_BYTES:
                    raise _structure_limit()
                if worksheet and name.rsplit("}", 1)[-1] == "row":
                    row_start = None

            def text(value):
                if text_lengths:
                    text_lengths[-1] += len(value)
                    if text_lengths[-1] > 64 * 1024:
                        raise _structure_limit()

            with archive.open(info) as stream:
                _walk_xml(stream, start=start, end=end, text=text, budget=budget)
            if worksheet:
                dimensions[info.filename] = (source_rows, source_columns)
    return dimensions

def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")



def _csv_rows(text: str, dialect):
    # csv.reader builds a whole logical row before column slicing. Limit its
    # input record, including quoted multiline fields, without StringIO's full
    # UTF-32 copy or splitlines() allocating every line in a large CSV.
    class Lines:
        position = 0
        record_chars = 0

        def __iter__(self):
            return self

        def __next__(self):
            start = self.position
            if start >= len(text):
                raise StopIteration
            limit = min(len(text), start + MAX_CSV_RECORD_CHARS + 1)
            ends = [pos for char in ("\n", "\r") if (pos := text.find(char, start, limit)) != -1]
            end = min(ends) + 1 if ends else limit
            if text[end - 1:end + 1] == "\r\n":
                end += 1
            self.record_chars += end - start
            if self.record_chars > MAX_CSV_RECORD_CHARS:
                raise _structure_limit()
            self.position = end
            return text[start:end]

    lines = Lines()
    reader = csv.reader(lines, dialect)
    while True:
        lines.record_chars = 0
        try:
            yield next(reader)
        except StopIteration:
            return

def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    return str(value)[:MAX_CELL_CHARS]


def _rows(source) -> tuple[list[list[str]], bool]:
    rows, chars, truncated = [], 0, False
    for index, raw in enumerate(islice(source, MAX_ROWS + 1)):
        if index == MAX_ROWS:
            truncated = True
            break
        values = list(islice(iter(raw), MAX_COLUMNS + 1))
        if len(values) > MAX_COLUMNS:
            truncated = True
        row = [_cell(v) for v in values[:MAX_COLUMNS]]
        if any(v is not None and len(str(v)) > MAX_CELL_CHARS for v in values):
            truncated = True
        chars += sum(map(len, row))
        if chars > MAX_TEXT_CHARS:
            truncated = True
            break
        rows.append(row)
    # Remove trailing empty columns, not interior empty cells (Excel coordinates matter).
    width = max((max((i + 1 for i, cell in enumerate(r) if cell), default=0) for r in rows), default=0)
    while rows and not any(rows[-1]):
        rows.pop()
    return [r[:width] + [""] * max(0, width - len(r)) for r in rows], truncated


def _table(names: list[str], sheet: int, source, *, row_count: int | None = None, column_count: int | None = None) -> dict:
    rows, truncated = _rows(source)
    return {"kind": "table", "sheets": names, "sheet": sheet, "rows": rows,
            "row_count": row_count, "column_count": column_count,
            "truncated": truncated or bool(row_count and row_count > MAX_ROWS) or bool(column_count and column_count > MAX_COLUMNS),
            "notice": "Предпросмотр только для чтения. Формулы не вычисляются, макросы не запускаются. Для работы с калькулятором скачайте оригинал и откройте в Excel или LibreOffice."}


def build_preview(data: bytes, filename: str, *, sheet: int = 0) -> dict:
    ext = pathlib.PurePath(filename).suffix.lower().lstrip(".")
    if ext not in SUPPORTED:
        raise HTTPException(415, "Для этого формата нет встроенного просмотра. Скачайте файл и откройте в подходящей программе.")
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "Файл больше 25 МБ. Для просмотра скачайте оригинал.")
    try:
        if ext in {"xlsx", "xlsm", "docx", "pptx"}:
            _check_zip(data)
        if ext in {"xlsx", "xlsm"}:
            dimensions = _check_workbook_structure(data)
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False, keep_vba=False)
            try:
                names = workbook.sheetnames
                if not 0 <= sheet < len(names):
                    raise HTTPException(422, "Лист не найден")
                ws = workbook[names[sheet]]
                rows, cols = dimensions.get(str(getattr(ws, "_worksheet_path", "")).lstrip("/"), (0, 0))
                return _table(names, sheet, ws.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLUMNS, values_only=True),
                              row_count=max(ws.max_row or 0, rows), column_count=max(ws.max_column or 0, cols))
            finally:
                workbook.close()
        if ext == "xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True, ragged_rows=True)
            try:
                if not 0 <= sheet < workbook.nsheets:
                    raise HTTPException(422, "Лист не найден")
                ws = workbook.sheet_by_index(sheet)
                source = (ws.row_values(i, 0, min(ws.ncols, MAX_COLUMNS)) for i in range(min(ws.nrows, MAX_ROWS + 1)))
                return _table(workbook.sheet_names(), sheet, source, row_count=ws.nrows, column_count=ws.ncols)
            finally:
                workbook.release_resources()
        if ext in {"csv", "tsv"}:
            text = _decode(data)
            if ext == "tsv":
                dialect = csv.excel_tab
            else:
                try:
                    dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
            return _table(["Таблица"], 0, _csv_rows(text, dialect))
        if ext in {"docx", "pptx"}:
            blocks, chars, truncated = [], 0, False
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if ext == "docx":
                    paths = ["word/document.xml"]
                    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                else:
                    paths = sorted([n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)],
                                   key=lambda n: int(re.search(r"slide(\d+)", n).group(1)))
                    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                    truncated = len(paths) > 100
                document_budget = {"nodes": 0}
                for i, path in enumerate(paths[:100]):
                    if ext == "pptx":
                        blocks.append({"type": "heading", "text": f"Слайд {i + 1}"})
                    paragraph, in_text, parts, part_chars = False, False, [], 0
                    tag_ns = ns.lstrip("{")

                    def start(name, attrs, position):
                        nonlocal paragraph, in_text, parts, part_chars
                        if name == tag_ns + "p":
                            paragraph, parts, part_chars = True, [], 0
                        if paragraph and name == tag_ns + "t":
                            in_text = True

                    def text(value):
                        nonlocal chars, part_chars, truncated
                        if not in_text:
                            return
                        remaining = MAX_TEXT_CHARS - chars - part_chars
                        if len(value) > remaining:
                            if remaining > 0:
                                parts.append(value[:remaining])
                                part_chars += remaining
                            if parts:
                                blocks.append({"type": "paragraph", "text": "".join(parts)})
                            truncated = True
                            raise _PreviewComplete()
                        parts.append(value)
                        part_chars += len(value)

                    def end(name, position):
                        nonlocal paragraph, in_text, chars, truncated
                        if name == tag_ns + "t":
                            in_text = False
                        if name == tag_ns + "p":
                            paragraph = False
                            value = "".join(parts)
                            if value.strip():
                                if len(blocks) >= 3000:
                                    truncated = True
                                    raise _PreviewComplete()
                                blocks.append({"type": "paragraph", "text": value})
                                chars += len(value)

                    try:
                        with archive.open(path) as stream:
                            _walk_xml(stream, start=start, end=end, text=text, budget=document_budget)
                    except _PreviewComplete:
                        break
                    if chars >= MAX_TEXT_CHARS or len(blocks) >= 3000:
                        truncated = True
                        break
            return {"kind": "document", "blocks": blocks, "truncated": truncated,
                    "notice": "Текстовый предпросмотр без исходной вёрстки, изображений и встроенных объектов. Полное оформление доступно в оригинале."}
        text = _decode(data)
        return {"kind": "text", "text": text[:MAX_TEXT_CHARS], "truncated": len(text) > MAX_TEXT_CHARS}
    except HTTPException:
        raise
    except Exception as exc:
        # Do not leak parser tracebacks, S3 credentials, or document contents.
        raise HTTPException(422, "Не удалось прочитать документ: формат повреждён, не поддерживается или защищён паролем. Скачайте оригинал.") from exc
