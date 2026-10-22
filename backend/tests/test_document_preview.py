import io
import zipfile

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import select

from app.models.asset import ContentAsset
from app.models.module import Module
from app.models.submodule_asset import SubmoduleAssetMap
from app.models.tag import Tag, ModuleTagMap, UserTagMap
from app.services.document_preview import build_preview


def xlsx_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Расчёт"
    sheet.append(["Позиция", "Цена", "Формула"])
    sheet.append(["Деталь", 120, "=B2*2"])
    workbook.create_sheet("Пояснения").append(["Только просмотр"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_xlsx_multisheet_formula_is_text_not_executed():
    preview = build_preview(xlsx_bytes(), "Калькулятор сталь._1743202857175.xlsx")
    assert preview["kind"] == "table"
    assert preview["sheets"] == ["Расчёт", "Пояснения"]
    assert preview["rows"][1] == ["Деталь", "120", "=B2*2"]
    assert not preview["truncated"]
    assert build_preview(xlsx_bytes(), "Расчёт.xlsx", sheet=1)["rows"] == [["Только просмотр"]]
    with pytest.raises(HTTPException) as error:
        build_preview(xlsx_bytes(), "Расчёт.xlsx", sheet=2)
    assert error.value.status_code == 422


def test_csv_and_tsv_cyrillic_encodings():
    data = "Название;Количество\nСталь;15".encode("cp1251")
    assert build_preview(data, "смета.csv")["rows"] == [["Название", "Количество"], ["Сталь", "15"]]
    assert build_preview("Цена\t100\nМасса\t20".encode("utf-16"), "таблица.tsv")["rows"][1] == ["Масса", "20"]


def test_preview_row_column_and_cell_limits(monkeypatch):
    from app.services import document_preview as preview
    monkeypatch.setattr(preview, "MAX_ROWS", 2)
    monkeypatch.setattr(preview, "MAX_COLUMNS", 2)
    monkeypatch.setattr(preview, "MAX_CELL_CHARS", 5)
    result = build_preview(b"abcdef,2,3\nsecond,4,5\nthird,6,7", "limits.csv")
    assert result["truncated"] is True
    assert len(result["rows"]) == 2
    assert all(len(row) == 2 for row in result["rows"])
    assert result["rows"][0][0] == "abcde"


def test_docx_text_preview_and_xml_entity_rejection():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Инструкция</w:t></w:r></w:p></w:body></w:document>')
    assert build_preview(stream.getvalue(), "файл.docx")["blocks"][0]["text"] == "Инструкция"
    malicious = io.BytesIO()
    with zipfile.ZipFile(malicious, "w") as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE a [<!ENTITY attack SYSTEM "file:///etc/passwd">]><a>&attack;</a>')
    with pytest.raises(HTTPException) as error:
        build_preview(malicious.getvalue(), "untrusted.docx")
    assert error.value.status_code == 422


@pytest.mark.parametrize("name,data,code", [("a.xlsx", b"broken", 422), ("a.html", b"<script>alert(1)</script>", 415), ("a.svg", b"<svg/>", 415), ("a.doc", b"binary", 415)])
def test_corrupt_and_unsafe_formats_have_clear_fallback(name, data, code):
    with pytest.raises(HTTPException) as error:
        build_preview(data, name)
    assert error.value.status_code == code


def test_zip_expansion_limit(monkeypatch):
    from app.services import document_preview
    monkeypatch.setattr(document_preview, "MAX_EXPANDED_BYTES", 10)
    with pytest.raises(HTTPException) as error:
        build_preview(xlsx_bytes(), "a.xlsx")
    assert error.value.status_code == 413


def attach(db, module, lesson, storage, *, name="Таблица.xlsx", content=None):
    asset = ContentAsset(module_id=module.id, bucket="testing", object_key=f"modules/{module.id}/file", original_filename=name)
    db.add(asset)
    db.flush()
    if lesson:
        db.add(SubmoduleAssetMap(submodule_id=lesson.id, asset_id=asset.id, order=1))
    db.commit()
    storage.objects[asset.object_key] = xlsx_bytes() if content is None else content
    return asset


def test_preview_authenticated_no_public_office_url_and_closes_body(client, db, make_module, make_user, storage):
    module, lesson, _ = make_module()
    _, headers = make_user()
    asset = attach(db, module, lesson, storage)
    response = client.get(f"/assets/{asset.id}/preview", headers=headers)
    assert response.status_code == 200, response.text
    assert response.json()["sheets"][0] == "Расчёт"
    assert "storage.test" not in response.text and "officeapps" not in response.text
    assert "no-store" in response.headers["cache-control"]
    assert all(body.closed for body in storage.bodies)
    assert client.get(f"/assets/{asset.id}/preview").status_code == 401


@pytest.mark.parametrize("state", ["draft", "restricted", "hidden", "archived_module", "archived_lesson", "archived_asset"])
def test_all_file_and_learning_routes_enforce_same_visibility(client, db, admin_headers, make_module, make_user, storage, state):
    module, lesson, quiz = make_module(required=True)
    _, headers = make_user()
    asset = attach(db, module, lesson, storage)
    if state == "draft":
        module.is_active = False
    elif state in {"restricted", "hidden"}:
        module.visibility = state
    elif state == "archived_module":
        module.is_archived = True
    elif state == "archived_lesson":
        lesson.is_archived = True
    else:
        asset.is_archived = True
    db.commit()
    for suffix in ["preview", "stream", "download", "presign-download"]:
        response = client.get(f"/assets/{asset.id}/{suffix}", headers=headers)
        assert response.status_code == 404, (suffix, state, response.text)
    assert not storage.bodies  # ACL is checked before touching private bytes.
    if state != "archived_asset":
        assert client.get(f"/submodules/{lesson.id}", headers=headers).status_code == 404
        assert client.post(f"/quizzes/{quiz.id}/start", headers=headers).status_code in {404, 409}
    assert client.get(f"/assets/{asset.id}/preview", headers=admin_headers).status_code == 200


def test_matching_tag_grants_access_not_just_catalog_entry(client, db, make_module, make_user, storage):
    module, lesson, _ = make_module()
    learner, headers = make_user()
    asset = attach(db, module, lesson, storage)
    module.visibility = "restricted"
    tag = Tag(name="Расчётчики")
    db.add(tag)
    db.flush()
    db.add_all([ModuleTagMap(module_id=module.id, tag_id=tag.id), UserTagMap(user_id=learner.id, tag_id=tag.id)])
    db.commit()
    assert client.get(f"/modules/{module.id}", headers=headers).status_code == 200
    assert client.get(f"/submodules/{lesson.id}", headers=headers).status_code == 200
    assert client.get(f"/assets/{asset.id}/preview", headers=headers).status_code == 200


def test_pdf_ranges_and_dangerous_inline_mime(client, db, make_module, make_user, storage):
    module, lesson, _ = make_module()
    _, headers = make_user()
    asset = attach(db, module, lesson, storage, name="Инструкция.pdf", content=b"%PDF-1.7 abcdefghijk")
    response = client.get(f"/assets/{asset.id}/stream", headers={**headers, "Range": "bytes=0-3"})
    assert response.status_code == 206, response.text
    assert response.content == b"%PDF"
    assert response.headers["x-frame-options"] == "SAMEORIGIN"
    assert response.headers["content-range"].startswith("bytes 0-3/")
    assert response.headers["content-type"] == "application/pdf"
    assert "sandbox" not in response.headers["content-security-policy"]
    assert client.get(f"/assets/{asset.id}/stream", headers={**headers, "Range": "bytes=999-"}).status_code == 416
    asset.original_filename = "danger.html"
    asset.mime_type = "text/html"
    db.commit()
    assert client.get(f"/assets/{asset.id}/stream", headers=headers).status_code == 415
    download = client.get(f"/assets/{asset.id}/download", headers=headers)
    assert download.status_code == 200
    assert download.headers["content-disposition"].startswith("attachment")
    assert all(body.closed for body in storage.bodies)


def rewrite_zip(data, replacements):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename in replacements:
                content = replacements[item.filename](content)
            target.writestr(item.filename, content)
    return output.getvalue()


def test_wrong_xlsx_dimension_does_not_hide_columns_or_rows():
    import re
    data = rewrite_zip(xlsx_bytes(), {"xl/worksheets/sheet1.xml": lambda value: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', value)})
    result = build_preview(data, "incorrect-dimension.xlsx")
    assert result["rows"][1] == ["Деталь", "120", "=B2*2"]
    assert (result["row_count"], result["column_count"]) == (2, 3)


def test_workbook_metadata_and_row_budgets_are_checked_before_openpyxl(monkeypatch):
    from app.services import document_preview as preview
    import openpyxl
    data = xlsx_bytes()
    def must_not_load(*args, **kwargs):
        raise AssertionError("unsafe input reached openpyxl")
    monkeypatch.setattr(openpyxl, "load_workbook", must_not_load)
    monkeypatch.setattr(preview, "MAX_XML_METADATA_BYTES", 30)
    with pytest.raises(HTTPException) as error:
        build_preview(data, "metadata.xlsx")
    assert error.value.status_code == 413
    monkeypatch.setattr(preview, "MAX_XML_METADATA_BYTES", 2 * 1024 * 1024)
    monkeypatch.setattr(preview, "MAX_XML_ROW_BYTES", 30)
    with pytest.raises(HTTPException) as error:
        build_preview(data, "oversized-row.xlsx")
    assert error.value.status_code == 413


def test_xml_parts_cannot_bypass_preflight_using_binary_names():
    data = rewrite_zip(xlsx_bytes(), {"[Content_Types].xml": lambda value: value.replace(b"/xl/worksheets/sheet1.xml", b"/xl/worksheets/unsafe.bin")})
    with pytest.raises(HTTPException) as error:
        build_preview(data, "unsafe.xlsx")
    assert error.value.status_code == 413


def test_docx_streams_without_full_tree_or_archive_read(monkeypatch):
    from app.services import document_preview as preview
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>' + "Учебный текст " * 40000 + '</w:t></w:r></w:p></w:body></w:document>')
    monkeypatch.setattr(preview, "MAX_TEXT_CHARS", 1000)
    def must_not_read(*args, **kwargs):
        raise AssertionError("whole ZIP entry read into memory")
    monkeypatch.setattr(zipfile.ZipFile, "read", must_not_read)
    result = build_preview(stream.getvalue(), "long.docx")
    assert result["truncated"] and sum(len(block["text"]) for block in result["blocks"]) == 1000


def test_xml_depth_and_unfinished_token_are_bounded():
    for content in ("<root>" + "<a>" * 70 + "x" + "</a>" * 70 + "</root>", '<root value="' + "x" * 300000 + '">text</root>'):
        data = io.BytesIO()
        with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", content)
        with pytest.raises(HTTPException) as error:
            build_preview(data.getvalue(), "bounded.docx")
        assert error.value.status_code == 413


def test_csv_quotes_multiline_crlf_and_column_flood_limit(monkeypatch):
    from app.services import document_preview as preview
    assert build_preview(b'a,"first\nsecond"\r\nb,third\r\n', "quotes.csv")["rows"] == [["a", "first\nsecond"], ["b", "third"]]
    monkeypatch.setattr(preview, "MAX_CSV_RECORD_CHARS", 100)
    with pytest.raises(HTTPException) as error:
        build_preview(b"," * 200, "too-wide.csv")
    assert error.value.status_code == 413


def test_preview_concurrency_limit_precedes_storage_reads(client, db, make_module, make_user, storage):
    from app.services.document_preview import preview_slot
    module, lesson, _ = make_module()
    _, headers = make_user()
    asset = attach(db, module, lesson, storage)
    with preview_slot(), preview_slot():
        result = client.get(f"/assets/{asset.id}/preview", headers=headers)
        assert result.status_code == 429 and result.headers["retry-after"] == "2"
        assert not storage.bodies
    assert client.get(f"/assets/{asset.id}/preview", headers=headers).status_code == 200


def test_pptx_slides_share_a_single_node_budget(monkeypatch):
    from app.services import document_preview as preview
    data = io.BytesIO()
    slide = '<root xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:p><a:r><a:t>Text</a:t></a:r></a:p></root>'
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/slides/slide1.xml", slide)
        archive.writestr("ppt/slides/slide2.xml", slide)
    monkeypatch.setattr(preview, "MAX_XML_NODES", 6)
    with pytest.raises(HTTPException) as error:
        build_preview(data.getvalue(), "slides.pptx")
    assert error.value.status_code == 413


def test_workbook_metadata_node_flood_is_rejected():
    data = rewrite_zip(xlsx_bytes(), {"xl/styles.xml": lambda value: b"<styleSheet>" + b"<x/>" * 50001 + b"</styleSheet>"})
    with pytest.raises(HTTPException) as error:
        build_preview(data, "styles.xlsx")
    assert error.value.status_code == 413
